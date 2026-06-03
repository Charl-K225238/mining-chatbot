"""
Construction du corpus textuel à partir des tables DuckDB et des fichiers uploadés.

Chaque document est un dict :
    {
        "id"      : str          identifiant unique
        "table"   : str          nom de la table source
        "date"    : str | None   date ISO si disponible
        "text"    : str          texte indexable
        "meta"    : dict         champs supplémentaires pour l'affichage
    }

Sources indexées :
    - descente_minerai        → OBSERVATION (pannes, incidents, commentaires)
    - journal                 → Commentaire (journal d'exploitation)
    - pges_actions            → Actions PGES + Sous actions + Livrable
    - suivi_actions           → Titre + Description + Observations
    - carburant_activites     → Observations
    - data/clean/txt/*.txt    → fichiers texte uploadés
    - data/clean/pdf/*.pdf    → documents PDF uploadés
    - data/clean/doc/*.docx   → documents Word uploadés
"""

import logging
from pathlib import Path
from typing import Any

import duckdb

logger = logging.getLogger(__name__)

_CLEAN_DIR = Path(__file__).parents[2] / "data" / "clean"


def build_corpus(con: duckdb.DuckDBPyConnection) -> list[dict[str, Any]]:
    """Extrait tous les documents textuels depuis DuckDB et les fichiers uploadés."""
    docs: list[dict[str, Any]] = []
    docs.extend(_from_descente_minerai(con))
    docs.extend(_from_journal(con))
    docs.extend(_from_pges(con))
    docs.extend(_from_suivi_actions(con))
    docs.extend(_from_carburant(con))
    docs.extend(_from_clean_files())
    logger.info("Corpus : %d documents extraits.", len(docs))
    return docs


# ── Extracteurs par table ────────────────────────────────────────────────────

def _from_descente_minerai(con: duckdb.DuckDBPyConnection) -> list[dict]:
    rows = con.execute("""
        SELECT
            DATE::VARCHAR          AS date,
            "Site"                 AS site,
            "Code / Nom engin / Immatriculation" AS engin,
            "QUANTITE JOURNALIERE DESCENDUE (T)" AS qty,
            "OBSERVATION"          AS obs
        FROM descente_minerai
        WHERE "OBSERVATION" IS NOT NULL
          AND LENGTH(TRIM("OBSERVATION")) > 2
        ORDER BY DATE
    """).fetchall()

    docs = []
    for i, (date, site, engin, qty, obs) in enumerate(rows):
        text = f"Observation mine : {obs}"
        if engin:
            text += f". Engin : {engin}"
        if qty:
            text += f". Tonnage : {qty:.0f} T"
        docs.append({
            "id":    f"dm_obs_{i}",
            "table": "descente_minerai",
            "date":  date,
            "text":  text,
            "meta":  {"site": site, "engin": engin, "qty": qty, "obs": obs, "date": date},
        })
    return docs


def _from_journal(con: duckdb.DuckDBPyConnection) -> list[dict]:
    try:
        rows = con.execute("""
            SELECT
                "Date"::VARCHAR  AS date,
                "Axe"            AS axe,
                "Objet"          AS objet,
                "Equipement"     AS equip,
                "Commentaire"    AS commentaire
            FROM journal
            WHERE "Commentaire" IS NOT NULL
              AND LENGTH(TRIM("Commentaire")) > 2
            ORDER BY "Date"
        """).fetchall()
    except Exception as e:
        logger.warning("journal non disponible : %s", e)
        return []

    docs = []
    for i, (date, axe, objet, equip, comment) in enumerate(rows):
        parts = [p for p in [axe, objet, equip, comment] if p]
        text = "Journal exploitation : " + " | ".join(parts)
        docs.append({
            "id":    f"jrn_{i}",
            "table": "journal",
            "date":  date,
            "text":  text,
            "meta":  {"axe": axe, "objet": objet, "equip": equip,
                      "commentaire": comment, "date": date},
        })
    return docs


def _from_pges(con: duckdb.DuckDBPyConnection) -> list[dict]:
    try:
        rows = con.execute("""
            SELECT
                "N°"                                       AS num,
                "Actions PGES / Recommandation ANDE"       AS action,
                "Sous actions"                             AS sous_action,
                "Livrable"                                 AS livrable,
                "Responsable"                              AS responsable,
                "Niveau d'avancement"                      AS avancement,
                "Commentaires"                             AS commentaire
            FROM pges_actions
            WHERE "N°" IS NOT NULL
        """).fetchall()
    except Exception as e:
        logger.warning("pges_actions non disponible : %s", e)
        return []

    docs = []
    for i, row in enumerate(rows):
        num, action, sous_action, livrable, resp, avancement, comment = row
        parts = [p for p in [action, sous_action, livrable, resp, avancement, comment] if p]
        text = f"Action PGES n°{num} : " + " | ".join(parts)
        docs.append({
            "id":    f"pges_{i}",
            "table": "pges_actions",
            "date":  None,
            "text":  text,
            "meta":  {"num": num, "action": action, "responsable": resp,
                      "avancement": avancement},
        })
    return docs


def _from_suivi_actions(con: duckdb.DuckDBPyConnection) -> list[dict]:
    try:
        rows = con.execute("""
            SELECT
                "Type"                  AS type_,
                "Titre des obligations" AS titre,
                "Description"           AS description,
                "Statuts"               AS statut,
                "Priorité"              AS priorite,
                "Observations"          AS obs
            FROM suivi_actions
            WHERE "Titre des obligations" IS NOT NULL
        """).fetchall()
    except Exception as e:
        logger.warning("suivi_actions non disponible : %s", e)
        return []

    docs = []
    for i, (type_, titre, desc, statut, prio, obs) in enumerate(rows):
        parts = [str(p) for p in [type_, titre, desc, statut, prio, obs] if p is not None]
        text = "Suivi obligation : " + " | ".join(parts)
        docs.append({
            "id":    f"suivi_{i}",
            "table": "suivi_actions",
            "date":  None,
            "text":  text,
            "meta":  {"titre": titre, "statut": statut, "priorite": prio},
        })
    return docs


def _from_carburant(con: duckdb.DuckDBPyConnection) -> list[dict]:
    try:
        rows = con.execute("""
            SELECT
                "Date"::VARCHAR   AS date,
                "MACHINE"         AS machine,
                "Observations"    AS obs
            FROM carburant_activites
            WHERE "Observations" IS NOT NULL
              AND LENGTH(TRIM("Observations")) > 2
            ORDER BY "Date"
        """).fetchall()
    except Exception as e:
        logger.warning("carburant_activites non disponible : %s", e)
        return []

    docs = []
    for i, (date, machine, obs) in enumerate(rows):
        text = f"Carburant - {machine or 'machine'} : {obs}"
        docs.append({
            "id":    f"carb_{i}",
            "table": "carburant_activites",
            "date":  date,
            "text":  text,
            "meta":  {"machine": machine, "obs": obs, "date": date},
        })
    return docs


_CHUNK_SIZE = 2000   # caractères max par chunk BM25


def _chunk_text(text: str, source_id: str, table: str,
                source_name: str) -> list[dict[str, Any]]:
    """Découpe un texte long en chunks de ~2000 chars pour un meilleur recall BM25."""
    chunks = []
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    current, chunk_i = [], 0
    cur_len = 0
    for para in paragraphs:
        if cur_len + len(para) > _CHUNK_SIZE and current:
            chunk_text = " ".join(current)
            chunks.append({
                "id":    f"{source_id}_c{chunk_i}",
                "table": table,
                "date":  None,
                "text":  chunk_text,
                "meta":  {"source": source_name, "chunk": chunk_i},
            })
            current, cur_len, chunk_i = [], 0, chunk_i + 1
        current.append(para)
        cur_len += len(para)
    if current:
        chunks.append({
            "id":    f"{source_id}_c{chunk_i}",
            "table": table,
            "date":  None,
            "text":  " ".join(current),
            "meta":  {"source": source_name, "chunk": chunk_i},
        })
    return chunks or [{"id": source_id, "table": table, "date": None,
                       "text": text[:_CHUNK_SIZE], "meta": {"source": source_name}}]


def _from_clean_files() -> list[dict[str, Any]]:
    """Extrait le texte des fichiers TXT, PDF et DOCX uploadés dans data/clean/."""
    docs: list[dict[str, Any]] = []

    # ── Fichiers TXT ─────────────────────────────────────────────────────────
    for path in sorted((_CLEAN_DIR / "txt").glob("*.txt")):
        try:
            text = path.read_text(encoding="utf-8", errors="replace").strip()
            if text:
                docs.extend(_chunk_text(text, f"txt_{path.stem}", "fichier_txt", path.name))
        except Exception as e:
            logger.warning("TXT %s : %s", path.name, e)

    # ── Fichiers PDF (1 chunk par page) ──────────────────────────────────────
    for path in sorted((_CLEAN_DIR / "pdf").glob("*.pdf")):
        try:
            import PyPDF2
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                pages_text = [(i + 1, (page.extract_text() or "").strip())
                              for i, page in enumerate(reader.pages)]
            nb_pages = len(pages_text)
            has_text = any(t for _, t in pages_text)
            if not has_text:
                logger.warning("PDF %s : aucun texte extrait (PDF scanné sans OCR ?)", path.name)
                continue
            for page_n, page_text in pages_text:
                if page_text:
                    docs.extend(_chunk_text(
                        page_text,
                        f"pdf_{path.stem}_p{page_n}",
                        "fichier_pdf",
                        f"{path.name} p.{page_n}/{nb_pages}",
                    ))
        except Exception as e:
            logger.warning("PDF %s : %s", path.name, e)

    # ── Fichiers DOCX ────────────────────────────────────────────────────────
    for path in sorted((_CLEAN_DIR / "doc").glob("*.docx")):
        try:
            from docx import Document
            doc = Document(path)
            # Paragraphes + texte des tableaux
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        parts.append(row_text)
            text = "\n".join(parts)
            if text:
                docs.extend(_chunk_text(text, f"docx_{path.stem}", "fichier_docx", path.name))
        except Exception as e:
            logger.warning("DOCX %s : %s", path.name, e)

    # ── Fichiers .doc — non indexables, avertissement ─────────────────────────
    doc_files = list((_CLEAN_DIR / "doc").glob("*.doc"))
    if doc_files:
        for f in doc_files:
            logger.warning(
                ".doc non indexé ('%s') — convertissez en .docx pour activer la recherche.",
                f.name
            )

    if docs:
        logger.info("Fichiers uploadés indexés : %d chunks.", len(docs))
    return docs
